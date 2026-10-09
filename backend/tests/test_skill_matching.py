"""Regression tests for shared CAMPUSLINK skill matching and resume parsing."""

import os
import tempfile
import unittest
from pathlib import Path

os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["CAMPUSLINK_SECRET_KEY"] = "test-only-campuslink-skill-matching-key"

from app import app, db, serializer
from ml.matching import match_student_to_job
from ml.skills import match_skill_lists, resolve_skill_key, split_skill_values
from models import Branch, Company, Course, Job, JobSkill, Skill, Student, User
from nlp.resume import analyze_resume, extract_text
from werkzeug.security import generate_password_hash


class SharedSkillMatchingTests(unittest.TestCase):
    def test_comma_separated_python_pandas_numpy_is_a_full_match(self):
        result = match_skill_lists("Python, Pandas, NumPy", "python, pandas, numpy")
        self.assertEqual(result["matched_skills"], ["Python", "Pandas", "NumPy"])
        self.assertEqual(result["missing_skills"], [])
        self.assertEqual(result["total_required_skills"], 3)
        self.assertEqual(result["match_percentage"], 100)

    def test_known_skills_split_when_entered_without_commas(self):
        self.assertEqual(split_skill_values("Python Pandas NumPy"), ["Python", "Pandas", "NumPy"])
        result = match_skill_lists("Python Pandas NumPy", "python pandas numpy")
        self.assertEqual(result["match_percentage"], 100)
        self.assertEqual(result["missing_skills"], [])

    def test_case_whitespace_aliases_arrays_and_database_records(self):
        result = match_skill_lists(
            [{"name": "  pOsTgReS  "}, {"skill": " jS "}],
            " PostgreSQL ; JavaScript ",
        )
        self.assertEqual(result["matched_skills"], ["PostgreSQL", "JavaScript"])
        self.assertEqual(result["match_percentage"], 100)
        self.assertEqual(resolve_skill_key("NumPy"), resolve_skill_key("numpy"))

    def test_partial_and_empty_lists_have_exact_percentages(self):
        partial = match_skill_lists("Python, SQL", "python; pandas; numpy")
        self.assertEqual(partial["matched_skills"], ["Python"])
        self.assertEqual(partial["missing_skills"], ["Pandas", "NumPy"])
        self.assertEqual(partial["total_required_skills"], 3)
        self.assertEqual(partial["match_percentage"], 33)
        empty = match_skill_lists(None, [])
        self.assertEqual(empty["matched_skills"], [])
        self.assertEqual(empty["missing_skills"], [])
        self.assertEqual(empty["total_required_skills"], 0)
        self.assertEqual(empty["match_percentage"], 0)

    def test_similar_unrelated_names_do_not_match(self):
        result = match_skill_lists(
            "Pythonista, Pandas-like, JavaScripted",
            ["Python", "Pandas", "JavaScript"],
        )
        self.assertEqual(result["matched_skills"], [])
        self.assertEqual(result["missing_skills"], ["Python", "Pandas", "JavaScript"])
        self.assertEqual(result["match_percentage"], 0)

    def test_recommendations_do_not_repeat_aliases_of_owned_skills(self):
        from ml.recommendations import rank_skill_candidates

        ranked, _ = rank_skill_candidates(
            {"skills": ["Postgres", "JS"], "career_goal": "Database Developer"},
            [{"skill": "PostgreSQL"}, {"skill": "JavaScript"}, {"skill": "Java"}],
        )
        recommended_keys = {resolve_skill_key(item["skill"]) for item in ranked}
        self.assertNotIn(resolve_skill_key("PostgreSQL"), recommended_keys)
        self.assertNotIn(resolve_skill_key("JavaScript"), recommended_keys)

    def test_legacy_space_separated_database_skill_row_is_split(self):
        result = match_skill_lists(
            ["Python Pandas NumPy"],
            ["python", "pandas", "numpy"],
            skill_catalog=["Python Pandas NumPy"],
        )
        self.assertEqual(result["match_percentage"], 100)
        self.assertEqual(result["total_required_skills"], 3)

    def test_match_function_uses_the_exact_required_skill_percentage(self):
        result = match_student_to_job(
            {"skills": "Python Pandas NumPy", "cgpa": 8, "attendance": 90, "semester": 6},
            {"minimum_cgpa": 0, "minimum_attendance": 0, "certifications": []},
            ["python", "pandas", "numpy"], [], [], [],
        )
        self.assertEqual(result["skill_match_percentage"], 100)
        self.assertEqual(result["components"]["skills"], 100)
        self.assertEqual(result["total_required_skills"], 3)

    def test_pdf_and_word_resume_skills_match_aliases(self):
        import fitz
        from docx import Document

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as folder:
            pdf_path = Path(folder) / "resume.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_textbox((50, 50, 500, 180), "TECHNICAL SKILLS\nPython, pandas; NumPy")
            pdf.save(pdf_path)

            docx_path = Path(folder) / "resume.docx"
            document = Document()
            document.add_paragraph("TECHNICAL SKILLS")
            document.add_paragraph("python  PANDAS  NumPy")
            document.save(docx_path)

            for path in (pdf_path, docx_path):
                with self.subTest(extension=path.suffix):
                    text = extract_text(path)
                    result = analyze_resume(
                        text, target_skills="python pandas numpy"
                    )
                    self.assertEqual(result["matched_skills"], ["Python", "Pandas", "NumPy"])
                    self.assertEqual(result["missing_skills"], [])
                    self.assertEqual(result["total_required_skills"], 3)
                    self.assertEqual(result["match_percentage"], 100)


class SkillMatchingIntegrationTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()
        with app.app_context():
            db.session.remove()
            db.drop_all()
            db.create_all()
            branch = Branch(name="CSE")
            student_user = User(name="Skill Test Student", email="skill-student@example.invalid",
                                password_hash=generate_password_hash("Password123!"), role="STUDENT")
            company_user = User(name="Skill Test Recruiter", email="skill-recruiter@example.invalid",
                                password_hash=generate_password_hash("Password123!"), role="COMPANY")
            db.session.add_all([branch, student_user, company_user])
            db.session.flush()
            student = Student(user_id=student_user.id, branch_id=branch.id, semester=6)
            company = Company(user_id=company_user.id, company_name="Skill Test Company")
            db.session.add_all([student, company])
            db.session.flush()
            job = Job(company_id=company.id, title="Python Data Role")
            required_skills = [Skill(name=name) for name in ("python", "pandas", "numpy")]
            db.session.add_all([job, *required_skills])
            db.session.flush()
            db.session.add_all([JobSkill(job_id=job.id, skill_id=skill.id) for skill in required_skills])
            db.session.commit()
            self.user_id = student_user.id
            self.job_id = job.id
            self.auth = {"Authorization": "Bearer " + serializer.dumps({"uid": student_user.id})}

    def tearDown(self):
        with app.app_context():
            db.session.remove()
            db.drop_all()

    def test_profile_save_and_job_match_endpoint_recognize_space_separated_skills(self):
        saved = self.client.put("/api/student/profile", json={
            "skills": "Python Pandas NumPy",
        }, headers=self.auth)
        self.assertEqual(saved.status_code, 200, saved.json)
        self.assertEqual([item["name"] for item in saved.json["profile"]["skills"]],
                         ["Python", "Pandas", "NumPy"])

        matched = self.client.post("/api/matching/calculate", json={
            "job_id": self.job_id,
        }, headers=self.auth)
        self.assertEqual(matched.status_code, 200, matched.json)
        result = matched.json["match"]
        self.assertEqual(result["matched_skills"], ["Python", "Pandas", "NumPy"])
        self.assertEqual(result["missing_skills"], [])
        self.assertEqual(result["total_required_skills"], 3)
        self.assertEqual(result["skill_match_percentage"], 100)

    def test_course_skill_filter_is_exact_and_alias_aware(self):
        with app.app_context():
            db.session.add_all([
                Course(title="PostgreSQL Guide", skill="PostgreSQL"),
                Course(title="JavaScript Guide", skill="JavaScript"),
                Course(title="Java Guide", skill="Java"),
            ])
            db.session.commit()

        js = self.client.get("/api/resources?skill=JS")
        self.assertEqual([item["title"] for item in js.json["items"]], ["JavaScript Guide"])
        java = self.client.get("/api/resources?skill=Java")
        self.assertEqual([item["title"] for item in java.json["items"]], ["Java Guide"])
        postgres = self.client.get("/api/resources?skill=Postgres")
        self.assertEqual([item["title"] for item in postgres.json["items"]], ["PostgreSQL Guide"])


if __name__ == "__main__":
    unittest.main()
