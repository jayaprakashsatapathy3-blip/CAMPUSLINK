import json
import os
import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError
from unittest.mock import patch

from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from werkzeug.security import generate_password_hash

os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["CAMPUSLINK_SECRET_KEY"] = "test-only-campuslink-features-signing-key"

from app import (GROQ_API_ENDPOINT, LEARNING_RESOURCE_YOUTUBE_URLS, AIServiceError, app, db,
                 branch_learning_focus,
                 adaptive_quiz_difficulty, ensure_schema_compatibility, full_quiz_serializer,
                 quick_quiz_local_date, quick_quiz_serializer, request_groq, request_quiz_generation, serializer)
from ai_service import OPENROUTER_CHAT_ENDPOINT
from nlp.intents import classify_intent, normalize_user_text
from nlp.resume import analyze_resume
from ml.prediction import placement_readiness
from ml.recommendations import rank_skill_candidates
from nlp.resume import extract_text as extract_resume_document
from models import (Branch, Company, Course, Job, JobSkill,
                    Notification, Placement, Quiz, QuizQuestion, QuickQuizAttempt,
                    QuizResult, Skill, Student, User)


class PlacementFeaturesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config.update(TESTING=True)
        with app.app_context():
            db.create_all()

    @classmethod
    def tearDownClass(cls):
        with app.app_context():
            db.session.remove()
            db.drop_all()

    def setUp(self):
        self.client = app.test_client()
        with app.app_context():
            db.session.remove()
            db.drop_all()
            db.create_all()
            self.cse = Branch(name="CSE")
            self.ece = Branch(name="ECE")
            db.session.add_all([self.cse, self.ece])
            db.session.flush()
            user = User(name="Feature Test Student", email="feature-student@example.invalid",
                        password_hash=generate_password_hash("Password123!"), role="STUDENT")
            db.session.add(user)
            db.session.flush()
            student = Student(user_id=user.id, student_number="TEST-001", branch_id=self.cse.id,
                              year=3, career_goal="Embedded Systems Engineer")
            db.session.add(student)
            db.session.commit()
            self.user_id = user.id
            self.student_id = student.id
            self.cse_id = self.cse.id
            self.ece_id = self.ece.id
            self.bearer = "Bearer " + serializer.dumps({"uid": user.id})

    def auth(self):
        return {"Authorization": self.bearer}

    def test_resources_replace_legacy_search_urls_with_direct_youtube_videos(self):
        for urls in LEARNING_RESOURCE_YOUTUBE_URLS.values():
            self.assertEqual(len(urls), 2)
            for url in urls:
                self.assertRegex(url, r"^https://www\.youtube\.com/watch\?v=[\w-]{11}$")

        with app.app_context():
            db.session.add_all([
                Course(title="Python Foundations", skill="Python", level="Beginner",
                       url="https://www.youtube.com/results?search_query=Python", duration="3–6 hours"),
                Course(title="Build a Practical Python Project", skill="Python", level="Intermediate",
                       url="https://www.youtube.com/results?search_query=Python+project", duration="Self-paced"),
                Course(title="Unmapped Resource", skill="Unknown Skill", level="Beginner", url="", duration="Self-paced"),
            ])
            db.session.commit()

        response = self.client.get("/api/resources?skill=Python")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json["items"]), 2)
        urls = {item["title"]: item["url"] for item in response.json["items"]}
        self.assertEqual(urls["Python Foundations"], "https://www.youtube.com/watch?v=rfscVS0vtbw")
        self.assertEqual(urls["Build a Practical Python Project"], "https://www.youtube.com/watch?v=qwAFL1597eM")
        self.assertTrue(all("youtube.com/watch?v=" in url for url in urls.values()))

        with app.app_context():
            persisted_urls = {course.title: course.url for course in Course.query.filter_by(skill="Python").all()}
            self.assertEqual(persisted_urls, urls)
            unmapped = Course.query.filter_by(skill="Unknown Skill").one()
            self.assertEqual(unmapped.url, "")

    def add_question_pool(self, categories, per_category=7, difficulty="Medium"):
        with app.app_context():
            for category in categories:
                quiz = Quiz(title=f"{category} Test Bank", category=category, active=True)
                db.session.add(quiz)
                db.session.flush()
                for index in range(per_category):
                    db.session.add(QuizQuestion(
                        quiz_id=quiz.id,
                        prompt=f"{category} test question {index + 1}?",
                        options=json.dumps(["Correct", "Incorrect one", "Incorrect two"]),
                        answer_index=0,
                        explanation=f"The correct answer is Correct for {category}.",
                        difficulty=difficulty,
                    ))
            db.session.commit()

    def test_branch_list_is_database_source_and_profile_branch_persists(self):
        response = self.client.get("/api/branches")
        self.assertEqual(response.status_code, 200)
        self.assertEqual({item["name"] for item in response.json["items"]}, {"CSE", "ECE"})

        saved = self.client.put("/api/student/profile", json={"branch": "  eCe  "}, headers=self.auth())
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json["profile"]["branch"], "ECE")
        self.assertEqual(saved.json["profile"]["branch_id"], self.ece_id)

        refreshed = self.client.get("/api/student/profile", headers=self.auth())
        self.assertEqual(refreshed.status_code, 200)
        self.assertEqual(refreshed.json["profile"]["branch"], "ECE")

        invalid = self.client.put("/api/student/profile", json={"branch": "Not a real branch"}, headers=self.auth())
        self.assertEqual(invalid.status_code, 400)
        still_saved = self.client.get("/api/student/profile", headers=self.auth())
        self.assertEqual(still_saved.json["profile"]["branch"], "ECE")

    def test_student_graduation_year_persists_validates_and_can_be_cleared(self):
        saved = self.client.put("/api/student/profile", json={"graduation_year": 2028}, headers=self.auth())
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json["profile"]["graduation_year"], 2028)

        refreshed = self.client.get("/api/student/profile", headers=self.auth())
        self.assertEqual(refreshed.json["profile"]["graduation_year"], 2028)

        for invalid_year in (2028.5, 2101):
            invalid = self.client.put(
                "/api/student/profile", json={"graduation_year": invalid_year}, headers=self.auth()
            )
            self.assertEqual(invalid.status_code, 400)

        cleared = self.client.put("/api/student/profile", json={"graduation_year": None}, headers=self.auth())
        self.assertEqual(cleared.status_code, 200)
        self.assertIsNone(cleared.json["profile"]["graduation_year"])

    def test_student_edit_profile_collections_and_gender_persist_and_can_be_removed(self):
        initial = self.client.put("/api/student/profile", json={
            "gender": "Prefer not to say",
            "courses": ["B.Tech Computer Science", "MCA", "b.tech computer science", ""],
            "achievements": ["Dean's List", "Hackathon finalist"],
            "interest_tags": ["Data analysis", "Python", "data analysis"],
        }, headers=self.auth())
        self.assertEqual(initial.status_code, 200, initial.json)
        profile = initial.json["profile"]
        self.assertEqual(profile["gender"], "Prefer not to say")
        self.assertEqual(profile["courses"], ["B.Tech Computer Science", "MCA"])
        self.assertEqual(profile["achievements"], ["Dean's List", "Hackathon finalist"])
        self.assertEqual(profile["interest_tags"], ["Data analysis", "Python"])
        self.assertEqual(profile["interests"], "Data analysis, Python")

        # A fresh GET represents reopening the editor or signing in again.
        reopened = self.client.get("/api/student/profile", headers=self.auth())
        self.assertEqual(reopened.status_code, 200)
        self.assertEqual(reopened.json["profile"]["courses"], profile["courses"])
        self.assertEqual(reopened.json["profile"]["achievements"], profile["achievements"])
        self.assertEqual(reopened.json["profile"]["interest_tags"], profile["interest_tags"])

        edited = self.client.put("/api/student/profile", json={
            "gender": "Other", "courses": ["M.Sc. Data Science"],
            "achievements": ["Published a research poster"], "interest_tags": ["Machine learning"],
        }, headers=self.auth())
        self.assertEqual(edited.status_code, 200, edited.json)
        self.assertEqual(edited.json["profile"]["gender"], "Other")
        self.assertEqual(edited.json["profile"]["courses"], ["M.Sc. Data Science"])
        self.assertEqual(edited.json["profile"]["achievements"], ["Published a research poster"])
        self.assertEqual(edited.json["profile"]["interest_tags"], ["Machine learning"])

        cleared = self.client.put("/api/student/profile", json={
            "courses": [], "achievements": [], "interest_tags": [],
        }, headers=self.auth())
        self.assertEqual(cleared.status_code, 200, cleared.json)
        self.assertEqual(cleared.json["profile"]["courses"], [])
        self.assertEqual(cleared.json["profile"]["achievements"], [])
        self.assertEqual(cleared.json["profile"]["interest_tags"], [])
        self.assertEqual(cleared.json["profile"]["interests"], "")

    def test_student_edit_profile_rejects_invalid_collection_data_without_mutation(self):
        saved = self.client.put("/api/student/profile", json={
            "gender": "Female", "courses": ["Computer Science"],
            "achievements": ["Dean's List"], "interest_tags": ["Python"],
        }, headers=self.auth())
        self.assertEqual(saved.status_code, 200)

        invalid_payloads = [
            {"gender": "Unknown"}, {"courses": "Computer Science"},
            {"achievements": [1]}, {"interest_tags": ["x" * 121]},
            {"courses": [f"Course {index}" for index in range(21)]},
        ]
        for payload in invalid_payloads:
            response = self.client.put("/api/student/profile", json=payload, headers=self.auth())
            self.assertEqual(response.status_code, 400, payload)

        unchanged = self.client.get("/api/student/profile", headers=self.auth()).json["profile"]
        self.assertEqual(unchanged["gender"], "Female")
        self.assertEqual(unchanged["courses"], ["Computer Science"])
        self.assertEqual(unchanged["achievements"], ["Dean's List"])
        self.assertEqual(unchanged["interest_tags"], ["Python"])

    def test_company_vacancy_accepts_typed_branches_and_normalizes_configured_case(self):
        with app.app_context():
            user = User(name="Feature Test Recruiter", email="feature-recruiter@example.invalid",
                        password_hash=generate_password_hash("Password123!"), role="COMPANY")
            db.session.add(user)
            db.session.flush()
            db.session.add(Company(user_id=user.id, company_name="Feature Test Company"))
            db.session.commit()
            company_token = "Bearer " + serializer.dumps({"uid": user.id})

        headers = {"Authorization": company_token}
        job_data = {"title": "ECE Developer", "description": "Build embedded software.",
                    "eligible_branches": ["ece"]}
        created = self.client.post("/api/company/jobs", json=job_data, headers=headers)
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.json["job"]["eligible_branches"], ["ECE"])

        job_data["title"] = "Unsupported Branch Role"
        job_data["eligible_branches"] = ["EEE"]
        custom_branch = self.client.post("/api/company/jobs", json=job_data, headers=headers)
        self.assertEqual(custom_branch.status_code, 201)
        self.assertEqual(custom_branch.json["job"]["eligible_branches"], ["EEE"])

    def test_company_accounts_are_isolated_across_jobs_applications_notifications_analytics_and_calendar(self):
        def create_company(name, email):
            with app.app_context():
                user = User(name=name, email=email,
                            password_hash=generate_password_hash("Password123!"), role="COMPANY")
                db.session.add(user)
                db.session.flush()
                company = Company(user_id=user.id, company_name=name)
                db.session.add(company)
                db.session.commit()
                return company.id, {"Authorization": "Bearer " + serializer.dumps({"uid": user.id})}

        first_company_id, first_auth = create_company("Isolation Company One", "isolation-one@example.invalid")
        second_company_id, second_auth = create_company("Isolation Company Two", "isolation-two@example.invalid")

        second_student_headers = None
        second_student_id = None
        with app.app_context():
            user = User(name="ECE Test Student", email="isolation-ece-student@example.invalid",
                        password_hash=generate_password_hash("Password123!"), role="STUDENT")
            db.session.add(user)
            db.session.flush()
            student = Student(user_id=user.id, student_number="ISOLATION-ECE", branch_id=self.ece_id,
                              year=4, career_goal="Embedded Engineer")
            db.session.add(student)
            db.session.commit()
            second_student_id = student.id
            second_student_headers = {"Authorization": "Bearer " + serializer.dumps({"uid": user.id})}

        first_job_response = self.client.post("/api/company/jobs", json={
            "title": "Python Role One", "eligible_branches": ["CSE"], "required_skills": ["Python"],
        }, headers=first_auth)
        second_job_response = self.client.post("/api/company/jobs", json={
            "title": "Java Role Two", "eligible_branches": ["ECE"], "required_skills": ["Java"],
        }, headers=second_auth)
        self.assertEqual(first_job_response.status_code, 201, first_job_response.json)
        self.assertEqual(second_job_response.status_code, 201, second_job_response.json)
        first_job_id = first_job_response.json["job"]["id"]
        second_job_id = second_job_response.json["job"]["id"]
        own_job_update = self.client.put(
            f"/api/company/jobs/{first_job_id}", json={"title": "Updated Python Role One"}, headers=first_auth
        )
        self.assertEqual(own_job_update.status_code, 200, own_job_update.json)

        first_application_response = self.client.post(
            f"/api/jobs/{first_job_id}/apply", json={}, headers=self.auth()
        )
        second_application_response = self.client.post(
            f"/api/jobs/{second_job_id}/apply", json={}, headers=second_student_headers
        )
        self.assertEqual(first_application_response.status_code, 201, first_application_response.json)
        self.assertEqual(second_application_response.status_code, 201, second_application_response.json)
        first_application_id = first_application_response.json["application"]["id"]
        second_application_id = second_application_response.json["application"]["id"]
        second_status_update = self.client.put(
            f"/api/applications/{second_application_id}/status", json={"status": "Shortlisted"}, headers=second_auth
        )
        self.assertEqual(second_status_update.status_code, 200, second_status_update.json)
        with app.app_context():
            db.session.add(Placement(student_id=second_student_id, company_id=second_company_id,
                                     job_id=second_job_id, package_lpa=12.5, placement_date=date.today()))
            first_notification = Notification(user_id=int(serializer.loads(first_auth["Authorization"][7:])["uid"]),
                                              title="Company One alert", message="One")
            second_notification = Notification(user_id=int(serializer.loads(second_auth["Authorization"][7:])["uid"]),
                                               title="Company Two alert", message="Two")
            db.session.add_all([first_notification, second_notification])
            db.session.commit()
            first_notification_id, second_notification_id = first_notification.id, second_notification.id

        first_event_response = self.client.post("/api/calendar", json={
            "title": "Company One interview day", "starts_at": "2026-11-01T10:00:00",
        }, headers=first_auth)
        second_event_response = self.client.post("/api/calendar", json={
            "title": "Company Two interview day", "starts_at": "2026-11-02T10:00:00",
        }, headers=second_auth)
        self.assertEqual(first_event_response.status_code, 201, first_event_response.json)
        self.assertEqual(second_event_response.status_code, 201, second_event_response.json)
        first_event_id = first_event_response.json["event"]["id"]
        second_event_id = second_event_response.json["event"]["id"]

        first_jobs = self.client.get("/api/company/jobs", headers=first_auth)
        second_jobs = self.client.get("/api/company/jobs", headers=second_auth)
        self.assertEqual([job["id"] for job in first_jobs.json["items"]], [first_job_id])
        self.assertEqual([job["id"] for job in second_jobs.json["items"]], [second_job_id])
        self.assertEqual([job["id"] for job in self.client.get("/api/jobs", headers=first_auth).json["items"]], [first_job_id])
        self.assertEqual([job["id"] for job in self.client.get("/api/jobs", headers=second_auth).json["items"]], [second_job_id])
        self.assertEqual(len(self.client.get("/api/jobs").json["items"]), 2)
        self.assertEqual(self.client.get(f"/api/jobs/{second_job_id}", headers=first_auth).status_code, 404)
        self.assertEqual(self.client.get(f"/api/jobs/{first_job_id}", headers=second_auth).status_code, 404)

        for method, path, kwargs in (
            ("put", f"/api/company/jobs/{second_job_id}", {"json": {"title": "Hijacked"}}),
            ("delete", f"/api/company/jobs/{second_job_id}", {}),
            ("get", f"/api/company/jobs/{second_job_id}/applications", {}),
            ("get", f"/api/company/jobs/{second_job_id}/candidates", {}),
            ("post", f"/api/company/jobs/{second_job_id}/candidates/{self.student_id}/shortlist", {"json": {}}),
            ("put", f"/api/applications/{second_application_id}/status", {"json": {"status": "Rejected"}}),
        ):
            response = getattr(self.client, method)(path, headers=first_auth, **kwargs)
            self.assertEqual(response.status_code, 404, (method, path, response.json))

        first_owned_applications = self.client.get(
            f"/api/company/jobs/{first_job_id}/applications", headers=first_auth
        )
        second_owned_applications = self.client.get(
            f"/api/company/jobs/{second_job_id}/applications", headers=second_auth
        )
        self.assertEqual([item["id"] for item in first_owned_applications.json["items"]], [first_application_id])
        self.assertEqual([item["id"] for item in second_owned_applications.json["items"]], [second_application_id])

        first_notifications = self.client.get("/api/notifications", headers=first_auth)
        second_notifications = self.client.get("/api/notifications", headers=second_auth)
        self.assertEqual([item["id"] for item in first_notifications.json["items"]], [first_notification_id])
        self.assertEqual([item["id"] for item in second_notifications.json["items"]], [second_notification_id])
        self.assertEqual(self.client.put(f"/api/notifications/{second_notification_id}/read", json={}, headers=first_auth).status_code, 404)

        first_events = self.client.get("/api/calendar", headers=first_auth)
        second_events = self.client.get("/api/calendar", headers=second_auth)
        self.assertEqual([item["id"] for item in first_events.json["items"]], [first_event_id])
        self.assertEqual([item["id"] for item in second_events.json["items"]], [second_event_id])
        self.assertEqual(self.client.put(f"/api/calendar/{second_event_id}", json={"title": "Hijacked"}, headers=first_auth).status_code, 404)
        self.assertEqual(self.client.delete(f"/api/calendar/{second_event_id}", headers=first_auth).status_code, 404)

        first_dashboard = self.client.get("/api/company/dashboard", headers=first_auth).json
        second_dashboard = self.client.get("/api/company/dashboard", headers=second_auth).json
        self.assertEqual([job["id"] for job in first_dashboard["jobs"]], [first_job_id])
        self.assertEqual([job["id"] for job in second_dashboard["jobs"]], [second_job_id])
        self.assertEqual(first_dashboard["metrics"]["applications"], 1)
        self.assertEqual(second_dashboard["metrics"]["applications"], 1)

        first_companies = self.client.get("/api/analytics/companies", headers=first_auth).json["items"]
        second_companies = self.client.get("/api/analytics/companies", headers=second_auth).json["items"]
        self.assertEqual([item["company"] for item in first_companies], ["Isolation Company One"])
        self.assertEqual([item["company"] for item in second_companies], ["Isolation Company Two"])
        first_skills = self.client.get("/api/analytics/skills", headers=first_auth).json["items"]
        second_skills = self.client.get("/api/analytics/skills", headers=second_auth).json["items"]
        self.assertEqual({item["skill"] for item in first_skills}, {"Python"})
        self.assertEqual({item["skill"] for item in second_skills}, {"Java"})
        filtered_first_skills = self.client.get(
            "/api/analytics/skills?company=Isolation%20Company%20Two", headers=first_auth
        ).json["items"]
        self.assertEqual({item["skill"] for item in filtered_first_skills}, {"Python"})
        first_placements = self.client.get("/api/analytics/placements", headers=first_auth).json
        second_placements = self.client.get("/api/analytics/placements", headers=second_auth).json
        self.assertEqual(first_placements["funnel"]["Applied"], 1)
        self.assertEqual(first_placements["packages"], [])
        self.assertEqual([row["branch"] for row in first_placements["branches"]], ["CSE"])
        self.assertEqual(second_placements["funnel"]["Applied"], 0)
        self.assertEqual(second_placements["funnel"]["Shortlisted"], 1)
        self.assertEqual(second_placements["packages"], [12.5])
        self.assertEqual([row["branch"] for row in second_placements["branches"]], ["ECE"])
        self.assertEqual(
            [row["branch"] for row in self.client.get("/api/analytics/branches", headers=first_auth).json["items"]],
            ["CSE"],
        )
        self.assertEqual(
            [row["branch"] for row in self.client.get("/api/analytics/branches", headers=second_auth).json["items"]],
            ["ECE"],
        )
        first_funnel = self.client.get("/api/analytics/applications", headers=first_auth).json["funnel"]
        second_funnel = self.client.get("/api/analytics/applications", headers=second_auth).json["funnel"]
        self.assertEqual(first_funnel["Applied"], 1)
        self.assertEqual(first_funnel["Shortlisted"], 0)
        self.assertEqual(second_funnel["Applied"], 0)
        self.assertEqual(second_funnel["Shortlisted"], 1)

    def test_branch_aware_chat_handles_custom_branches_and_typo_normalization(self):
        self.assertEqual(normalize_user_text("plz help me with an interveiw"), "please help me with an interview")
        self.assertEqual(classify_intent("plz help me with an interveiw")['intent'], "interview")
        self.client.put("/api/student/profile", json={"branch": "CSE"}, headers=self.auth())
        body = {"messages": [{"role": "user", "content": "How should I prepare for an interview?"}], "mode": "placement"}
        with patch("app.request_groq", return_value=("Start with one project example.", "groq-test")) as call:
            response = self.client.post("/api/chatbot", json=body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["branch"], "CSE")
        self.assertEqual(response.json["intent"], "interview")
        self.assertIn("programming, data structures", call.call_args.args[1])
        self.assertIn("stated specialization", branch_learning_focus("CSE AI"))

    def test_chat_short_followup_inherits_prior_intent(self):
        body = {"messages": [
            {"role": "user", "content": "Help me prepare for a technical interview."},
            {"role": "assistant", "content": "What role are you targeting?"},
            {"role": "user", "content": "CSE AI"},
        ]}
        with patch("app.request_groq", return_value=("Let's focus on your AI project.", "groq-test")) as call:
            response = self.client.post("/api/chatbot", json=body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertIn('"inherited_from_history":true', call.call_args.args[1])
        self.assertEqual(response.json["intent"], "interview")

    def test_resume_analysis_extracts_skills_projects_and_qualifications(self):
        result = analyze_resume("""EDUCATION
B.Tech Computer Science, 2024
PROJECTS
Campus job matching app built with Python and SQL.
SKILLS
Python, SQL
""", skill_catalog=["Python", "SQL", "React"], target_skills=["Python", "SQL", "React"])
        self.assertIn("Python", result["detected_skills"])
        self.assertEqual(result["qualifications"], ["B.Tech Computer Science, 2024"])
        self.assertIn("Campus job matching app built with Python and SQL.", result["projects"])
        self.assertEqual(result["missing_skills"], ["React"])
        self.assertEqual(result["matched_skills"], ["Python", "SQL"])
        self.assertEqual(result["match_percentage"], 67)
        self.assertEqual(result["missing_skills"], ["React"])

    def test_resume_analysis_reports_partial_matches_and_additional_skills(self):
        result = analyze_resume(
            "EDUCATION\nB.Tech Computer Science\nTECHNICAL SKILLS\nMySQL, Excel, Python\n",
            target_skills=["SQL", "Power BI"],
        )
        self.assertEqual(result["matched_skills"], [])
        self.assertEqual(result["partially_matched_skills"], [{"skill": "SQL", "evidence": ["MySQL"]}])
        self.assertEqual(result["missing_skills"], ["Power BI"])
        self.assertIn("Python", result["additional_relevant_skills"])
        self.assertIn("Excel", result["additional_relevant_skills"])

    def test_resume_docx_parser_reads_paragraphs_and_tables(self):
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "resume.docx"
            document = WordDocument()
            document.add_paragraph("EDUCATION")
            document.add_paragraph("B.Tech Computer Science")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Projects"
            table.cell(0, 1).text = "Python analytics project"
            document.save(path)
            extracted = extract_resume_document(path)
        self.assertIn("B.Tech Computer Science", extracted)
        self.assertIn("Python analytics project", extracted)

    def test_resume_doc_parser_uses_available_antiword_backend(self):
        from types import SimpleNamespace

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "legacy.doc"
            path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"legacy")
            with patch("nlp.resume.shutil.which", side_effect=lambda name: "antiword.exe" if name == "antiword" else None), \
                    patch("nlp.resume.subprocess.run", return_value=SimpleNamespace(
                        returncode=0, stdout=b"EDUCATION\nB.Tech\nPython and SQL project evidence")):
                extracted = extract_resume_document(path)
        self.assertIn("Python and SQL", extracted)

    def test_resume_docx_upload_analyzes_uploaded_job_description(self):
        from docx import Document as WordDocument

        document = WordDocument()
        document.add_paragraph("EDUCATION")
        document.add_paragraph("B.Tech Computer Science, 2025")
        document.add_paragraph("TECHNICAL SKILLS")
        document.add_paragraph("Python, SQL, Excel")
        resume_file = BytesIO()
        document.save(resume_file)
        resume_file.seek(0)
        ai_review = {"summary": "The resume has Python and SQL evidence but lacks Power BI.",
                     "strengths": [], "risks": [], "experience_feedback": "", "education_feedback": "",
                     "recommendations": []}
        original_folder = app.config["UPLOAD_FOLDER"]
        with tempfile.TemporaryDirectory() as folder:
            app.config["UPLOAD_FOLDER"] = folder
            try:
                with patch("app.request_groq", return_value=(json.dumps(ai_review), "groq-test")):
                    response = self.client.post("/api/resume/analyze", data={
                        "file": (resume_file, "student-resume.docx"),
                        "job_description_file": (BytesIO(b"Junior Data Analyst requires Python, SQL, Power BI, and Excel."), "data-analyst.txt"),
                    }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(response.status_code, 201, response.json)
                report = response.json["analysis"]
                self.assertEqual(report["target_role"], "Provided job description")
                self.assertEqual(report["job_description_filename"], "data-analyst.txt")
                self.assertEqual(report["matched_skills"], ["Python", "SQL", "Excel"])
                self.assertEqual(report["missing_skills"], ["Power BI"])
                self.assertEqual(report["match_percentage"], 75)
            finally:
                app.config["UPLOAD_FOLDER"] = original_folder

    def test_resume_upload_rejects_invalid_file_and_scanned_pdf_gracefully(self):
        import fitz
        from docx import Document as WordDocument

        original_folder = app.config["UPLOAD_FOLDER"]
        with tempfile.TemporaryDirectory() as folder:
            app.config["UPLOAD_FOLDER"] = folder
            try:
                invalid = self.client.post("/api/resume/analyze", data={
                    "file": (BytesIO(b"not a PDF"), "resume.pdf"),
                }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(invalid.status_code, 400)
                self.assertIn("not a valid PDF", invalid.json["error"])

                unsupported = self.client.post("/api/resume/analyze", data={
                    "file": (BytesIO(b"plain text"), "resume.rtf"),
                }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(unsupported.status_code, 400)
                self.assertIn("Unsupported file type", unsupported.json["error"])

                empty_word = BytesIO()
                WordDocument().save(empty_word)
                empty_word.seek(0)
                empty = self.client.post("/api/resume/analyze", data={
                    "file": (empty_word, "empty.docx"),
                }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(empty.status_code, 400)
                self.assertIn("No readable text", empty.json["error"])

                document = fitz.open()
                document.new_page()
                scanned = self.client.post("/api/resume/analyze", data={
                    "file": (BytesIO(document.tobytes()), "scanned.pdf"),
                }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(scanned.status_code, 400)
                self.assertIn("scanned", scanned.json["error"])
                self.assertIn("OCR is not configured", scanned.json["error"])
            finally:
                app.config["UPLOAD_FOLDER"] = original_folder

    def test_resume_upload_analyzes_pdf_against_selected_job_and_persists_ai_report(self):
        import fitz

        with app.app_context():
            recruiter = User(name="Resume Test Recruiter", email="resume-recruiter@example.invalid",
                             password_hash=generate_password_hash("Password123!"), role="COMPANY")
            db.session.add(recruiter)
            db.session.flush()
            company = Company(user_id=recruiter.id, company_name="Resume Test Analytics")
            db.session.add(company)
            db.session.flush()
            job = Job(company_id=company.id, title="Junior Data Analyst",
                      description="Analyze product datasets and communicate findings.")
            required = [Skill(name=name) for name in ("Python", "SQL", "Power BI")]
            db.session.add_all([job, *required])
            db.session.flush()
            db.session.add_all([JobSkill(job_id=job.id, skill_id=skill.id) for skill in required])
            job_id = job.id
            db.session.commit()

        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), """EDUCATION
B.Tech Computer Science, 2025
EXPERIENCE
Junior analyst cleaned product data with Python and MySQL.
PROJECTS
Built an Excel dashboard to summarize monthly sales trends.
CERTIFICATIONS
Analytics Foundations Certificate
SKILLS
Python, MySQL, Excel
candidate@example.com 9876543210
""")
        pdf = document.tobytes()
        ai_review = {
            "summary": "Python evidence is clear, while the resume shows MySQL rather than direct SQL project depth.",
            "strengths": ["A relevant analytics project is described."],
            "risks": ["Power BI is not evidenced."],
            "experience_feedback": "The analyst entry names Python and a data-cleaning responsibility.",
            "education_feedback": "A computer science degree is listed.",
            "recommendations": [{
                "skill": "Power BI", "priority": "High", "reason": "It is required by the selected vacancy.",
                "project": "Build a Power BI sales dashboard.", "exercise": "Create measures for monthly revenue and growth.",
            }],
        }
        original_folder = app.config["UPLOAD_FOLDER"]
        with tempfile.TemporaryDirectory() as folder:
            app.config["UPLOAD_FOLDER"] = folder
            try:
                with patch("app.request_groq", return_value=(json.dumps(ai_review), "groq-test")) as ai_call:
                    response = self.client.post("/api/resume/analyze", data={
                        "file": (BytesIO(pdf), "candidate-resume.pdf"),
                        "selected_job_id": str(job_id),
                    }, content_type="multipart/form-data", headers=self.auth())
                self.assertEqual(response.status_code, 201, response.json)
                report = response.json["analysis"]
                self.assertEqual(report["target_role"], "Junior Data Analyst")
                self.assertEqual(report["matched_skills"], ["Python"])
                self.assertEqual(report["partially_matched_skills"], [{"skill": "SQL", "evidence": ["MySQL"]}])
                self.assertEqual(report["missing_skills"], ["Power BI"])
                self.assertEqual(report["match_percentage"], 33)
                self.assertEqual(report["ai_review"]["status"], "ready")
                self.assertEqual(report["learning_roadmap"][0]["project"], "Build a Power BI sales dashboard.")
                sent_content = ai_call.call_args.args[0][0]["content"]
                self.assertIn("EDUCATION", sent_content)
                self.assertIn("[email removed]", sent_content)
                self.assertIn("[phone removed]", sent_content)
                self.assertNotIn("candidate@example.com", sent_content)
                self.assertNotIn("9876543210", sent_content)
                latest = self.client.get("/api/resume/analysis", headers=self.auth())
                self.assertEqual(latest.status_code, 200)
                self.assertEqual(latest.json["analysis"]["match_percentage"], 33)
                profile = self.client.get("/api/student/profile", headers=self.auth())
                self.assertEqual(profile.status_code, 200)
                stored_resume = profile.json["profile"]["resume_filename"]
                self.assertTrue(Path(folder, stored_resume).is_file())
                self.assertNotEqual(stored_resume, "candidate-resume.pdf")
            finally:
                app.config["UPLOAD_FOLDER"] = original_folder

    def test_skill_recommendation_ranking_explains_ml_and_profile_signals(self):
        ranked, method = rank_skill_candidates(
            {"branch": "CSE-AI", "career_goal": "Machine Learning Engineer", "skills": ["Python"],
             "weak_quiz_topics": ["Machine Learning"]},
            [{"skill": "Machine Learning", "vacancies": 4, "resource_titles": ["Machine Learning Basics"]},
             {"skill": "Power BI", "vacancies": 2, "resource_titles": ["Power BI Introduction"]}],
        )
        self.assertTrue(method.startswith("tfidf_"))
        self.assertEqual(ranked[0]["skill"], "Machine Learning")
        self.assertIn("quiz practice", ranked[0]["reason"])

    def test_skill_recommendations_prioritize_data_skills_and_remove_duplicates_or_irrelevant_items(self):
        ranked, _ = rank_skill_candidates(
            {"branch": "", "career_goal": "Data Analyst", "interests": "analytics, dashboards",
             "skills": ["Excel"]},
            [{"skill": "Python", "vacancies": 0}, {"skill": "python", "vacancies": 0},
             {"skill": "SQL", "vacancies": 0}, {"skill": "Power BI", "vacancies": 0},
             {"skill": "React", "vacancies": 0}, {"skill": "Java", "vacancies": 0},
             {"skill": "Data Structures", "vacancies": 0}],
        )
        names = [item["skill"].casefold() for item in ranked]
        self.assertEqual(len(names), len(set(names)))
        self.assertIn("python", names)
        self.assertIn("sql", names)
        self.assertIn("power bi", names)
        self.assertNotIn("react", names)
        self.assertNotIn("java", names)
        self.assertNotIn("data structures", names)
        self.assertTrue(any("career goal" in item["reason"] for item in ranked))
        self.assertTrue(any("interests" in item["reason"] for item in ranked))

    def test_skill_recommendations_skip_unrelated_candidates(self):
        ranked, _ = rank_skill_candidates(
            {"branch": "", "career_goal": "Graphic Designer", "interests": "landscape painting",
             "skills": ["Illustration"]},
            [{"skill": "Python", "vacancies": 0}, {"skill": "SQL", "vacancies": 0}],
        )
        self.assertEqual(ranked, [])

    def test_skill_recommendations_endpoint_uses_saved_interests_without_active_vacancies(self):
        with app.app_context():
            db.session.add_all([Skill(name=name) for name in (
                "Python", "SQL", "Power BI", "React", "Excel", "Java", "Data Structures",
            )])
            db.session.commit()

        saved = self.client.put("/api/student/profile", json={
            "career_goal": "Data Analyst", "interests": "analytics, dashboards", "skills": ["Excel"],
        }, headers=self.auth())
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json["profile"]["interests"], "analytics, dashboards")

        response = self.client.get("/api/recommendations/skills", headers=self.auth())
        self.assertEqual(response.status_code, 200)
        names = [item["skill"].casefold() for item in response.json["items"]]
        self.assertIn("python", names)
        self.assertIn("sql", names)
        self.assertNotIn("react", names)
        self.assertNotIn("java", names)
        self.assertNotIn("data structures", names)
        self.assertEqual(len(names), len(set(names)))
        self.assertEqual(response.json["based_on"]["interests"], "analytics, dashboards")
        self.assertTrue(all(item["vacancies"] == 0 for item in response.json["items"]))

    def test_readiness_model_uses_holdout_validation_and_excludes_target_student(self):
        with app.app_context():
            target = db.session.get(Student, self.student_id)
            target.year = 4
            for index in range(40):
                user = User(name=f"Cohort Student {index}", email=f"cohort-{index}@example.invalid",
                            password_hash=generate_password_hash("Password123!"), role="STUDENT")
                db.session.add(user)
                db.session.flush()
                record = Student(user_id=user.id, student_number=f"COHORT-{index}", branch_id=self.cse_id,
                                 year=4, cgpa=5.5 + (index % 6) * 0.6, attendance=65 + (index % 7) * 5)
                db.session.add(record)
                db.session.flush()
                if index < 15:
                    db.session.add(Placement(student_id=record.id))
            db.session.commit()
            result = placement_readiness(target)
        self.assertEqual(result["method"], "Logistic-regression readiness estimate")
        self.assertEqual(result["historical_records"], 40)
        self.assertEqual(result["validation"]["split"], "stratified 75/25 holdout")
        self.assertEqual(result["validation"]["test_records"], 10)
        self.assertIn("pre-placement feature snapshots are unavailable", result["note"])

    def test_readiness_uses_transparent_rule_when_validation_cohort_is_too_small(self):
        response = self.client.get("/api/student/dashboard", headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["metrics"]["readiness_method"], "Transparent rule-based readiness estimate")
        self.assertIsNone(response.json["metrics"]["readiness_validation"])

    def test_company_job_form_loads_branches_from_the_backend(self):
        with open(os.path.join(os.path.dirname(__file__), "../../frontend/assets/js/app.js"), encoding="utf-8") as source_file:
            javascript = source_file.read()
        self.assertIn("const {items:branchItems}=await api('/branches')", javascript)
        self.assertIn('value="${esc(branch)}"', javascript)

    def test_student_registration_requires_and_persists_college_name(self):
        payload = {
            "name": "New Campus Student",
            "email": "new-campus-student@example.invalid",
            "password": "Password123!",
            "role": "STUDENT",
            "phone": "1234567890",
        }
        missing = self.client.post("/api/auth/register", json=payload)
        self.assertEqual(missing.status_code, 400)
        self.assertIn("College name", missing.json["error"])

        payload["college_name"] = "North Campus College"
        registered = self.client.post("/api/auth/register", json=payload)
        self.assertEqual(registered.status_code, 201)
        headers = {"Authorization": f"Bearer {registered.json['token']}"}
        profile = self.client.get("/api/student/profile", headers=headers)
        self.assertEqual(profile.status_code, 200)
        self.assertEqual(profile.json["profile"]["college_name"], "North Campus College")

        updated = self.client.put(
            "/api/student/profile", json={"college_name": "Updated Campus"}, headers=headers
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json["profile"]["college_name"], "Updated Campus")

    def test_company_registration_requires_and_persists_company_name(self):
        payload = {
            "name": "Company Recruiter",
            "email": "company-recruiter@example.invalid",
            "password": "Password123!",
            "role": "COMPANY",
        }
        missing = self.client.post("/api/auth/register", json=payload)
        self.assertEqual(missing.status_code, 400)
        self.assertIn("Company name", missing.json["error"])

        payload["company_name"] = "Campus Systems Inc"
        registered = self.client.post("/api/auth/register", json=payload)
        self.assertEqual(registered.status_code, 201)
        self.assertIsNone(registered.json["user"]["phone"])
        headers = {"Authorization": f"Bearer {registered.json['token']}"}
        profile = self.client.get("/api/company/profile", headers=headers)
        self.assertEqual(profile.status_code, 200)
        self.assertEqual(profile.json["profile"]["company_name"], "Campus Systems Inc")

    def test_role_specific_registration_field_mapping_and_switch_reset(self):
        with open(os.path.join(os.path.dirname(__file__), "../../frontend/assets/js/app.js"), encoding="utf-8") as source_file:
            javascript = source_file.read()
        self.assertIn("name: 'college_name', label: 'College Name', placeholder: 'Enter your college name'", javascript)
        self.assertIn("name: 'company_name', label: 'Company Name', placeholder: 'Enter your company name'", javascript)
        self.assertIn("if (input.name !== nextName) input.value = '';", javascript)
        self.assertIn("syncOrganizationField(form,event.target.value)", javascript)
        self.assertNotIn("Organization name", javascript)

    def test_frontend_assets_include_global_assistant_and_dashboard_reminder_without_key(self):
        index = self.client.get("/")
        javascript = self.client.get("/assets/js/app.js")
        stylesheet = self.client.get("/assets/css/app.css")
        self.assertEqual(index.status_code, 200)
        self.assertEqual(javascript.status_code, 200)
        self.assertEqual(stylesheet.status_code, 200)
        self.assertIn(b"renderAssistantWidget", javascript.data)
        self.assertIn(b"Professional readiness", javascript.data)
        self.assertIn(b"readiness_validation", javascript.data)
        self.assertIn(b"Reminder", javascript.data)
        self.assertNotIn(b"GROQ_API_KEY", javascript.data)
        self.assertIn(b"assistant-panel", stylesheet.data)
        index.close()
        javascript.close()
        stylesheet.close()

    def test_groq_route_uses_branch_context_and_returns_service_output(self):
        self.client.put("/api/student/profile", json={"branch_id": self.ece_id}, headers=self.auth())
        request_body = {"messages": [{"role": "user", "content": "Help me prepare for placements."}],
                        "mode": "mock_interview"}
        with patch("app.request_groq", return_value=("What project are you most proud of?", "groq-test")) as call:
            response = self.client.post("/api/chatbot", json=request_body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["answer"], "What project are you most proud of?")
        self.assertEqual(response.json["branch"], "ECE")
        system_instruction = call.call_args.args[1]
        self.assertIn('"branch":"ECE"', system_instruction)
        self.assertIn("ask exactly one suitable question", system_instruction)

    def test_groq_route_requires_auth_and_surfaces_configuration_error(self):
        no_auth = self.client.post("/api/chatbot", json={"message": "Hello"})
        self.assertEqual(no_auth.status_code, 401)
        with patch("app.request_groq", side_effect=AIServiceError("Groq is not configured.")):
            response = self.client.post("/api/chatbot", json={"message": "Hello"}, headers=self.auth())
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json["error"], "Groq is not configured.")

    def test_communication_coach_uses_existing_groq_route_and_returns_coaching_metadata(self):
        response_json = {
            "reply": "Good attempt! Say: ‘I am studying CSE.’ We use ‘am studying’ for an activity happening now. What are you learning this week?",
            "difficulty": "intermediate",
            "evaluation": {"grammar": 6, "vocabulary": 7, "clarity": 8, "fluency": 7,
                           "relevance": 8, "professionalism": 7, "naturalness": 7, "overall": 7,
                           "note": "Your meaning was clear; keep practicing the present continuous."},
            "corrections": ["Use ‘am studying’ for an action happening now."],
            "corrected_version": "I am studying CSE.",
            "mistake_explanations": ["Use am + verb-ing for an activity happening now."],
            "better_response": "I am currently studying computer science.",
            "vocabulary_suggestions": ["currently — at the present time"],
            "vocabulary_learned": ["present continuous"],
        }
        request_body = {
            "mode": "communication_coach", "practice_mode": "daily", "difficulty": "beginner",
            "session_memory": {"turn_count": 1, "recent_answers": ["I am study in CSE."],
                               "recent_corrections": [], "learned_vocabulary": []},
            "messages": [{"role": "user", "content": "I am study in CSE."}],
        }
        with patch("app.request_groq", return_value=(json.dumps(response_json), "groq-test")) as call:
            response = self.client.post("/api/chatbot", json=request_body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertIn("I am studying CSE", response.json["answer"])
        self.assertEqual(response.json["difficulty"], "intermediate")
        self.assertEqual(response.json["evaluation"]["grammar"], 6)
        self.assertEqual(response.json["corrections"], response_json["corrections"])
        self.assertEqual(response.json["corrected_version"], response_json["corrected_version"])
        self.assertEqual(response.json["mistake_explanations"], response_json["mistake_explanations"])
        self.assertEqual(response.json["better_response"], response_json["better_response"])
        self.assertEqual(response.json["vocabulary_suggestions"], response_json["vocabulary_suggestions"])
        instruction = call.call_args.args[1]
        self.assertIn('"branch":"CSE"', instruction)
        self.assertIn('"year":3', instruction)
        self.assertIn('"career_goal":"Embedded Systems Engineer"', instruction)
        self.assertIn("Daily Conversation", instruction)
        self.assertIn("at most one question", instruction)
        self.assertIn("I am study in CSE.", instruction)

    def test_communication_report_returns_validated_rubric_and_rejects_empty_or_invalid_turns(self):
        report = {
            "overall": 8, "grammar": 7, "vocabulary": 8, "clarity": 8, "fluency": 7,
            "professional_communication": 8, "strengths": ["Clear explanation"],
            "common_mistakes": ["Past tense consistency"], "new_vocabulary": ["concise — brief and clear"],
            "improvement_areas": ["Use past tense consistently"],
            "recommendations": ["Practice a short project story", "Review irregular past-tense verbs"],
            "next_mode": "interview",
        }
        request_body = {
            "mode": "communication_report", "practice_mode": "interview", "difficulty": "intermediate",
            "session_memory": {"turn_count": 2, "recent_corrections": ["past tense"],
                               "learned_vocabulary": ["concise"]},
            "messages": [{"role": "user", "content": "Tell me about a project."},
                         {"role": "assistant", "content": "What did you learn?"},
                         {"role": "user", "content": "I learn to explain clearly."},
                         {"role": "user", "content": "Prepare my practice report."}],
        }
        with patch("app.request_groq", return_value=(json.dumps(report), "groq-test")) as call:
            response = self.client.post("/api/chatbot", json=request_body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["report"]["overall"], 8)
        self.assertEqual(response.json["report"]["next_mode"], "interview")
        self.assertEqual(len(response.json["report"]["recommendations"]), 2)
        self.assertIn("only evidence in the visible transcript", call.call_args.args[1])

        empty = self.client.post("/api/chatbot", json={"mode": "communication_coach", "practice_mode": "daily",
                                                       "messages": []}, headers=self.auth())
        self.assertEqual(empty.status_code, 400)
        invalid_mode = self.client.post("/api/chatbot", json={"mode": "communication_coach",
                                                              "practice_mode": "unlisted",
                                                              "messages": [{"role": "user", "content": "Hello"}]},
                                        headers=self.auth())
        self.assertEqual(invalid_mode.status_code, 400)

    def test_communication_coach_surfaces_ai_failure_without_fallback_reply(self):
        request_body = {"mode": "communication_coach", "practice_mode": "grammar", "difficulty": "beginner",
                        "messages": [{"role": "user", "content": "He go to class."}]}
        with patch("app.request_groq", side_effect=AIServiceError("Groq is temporarily unavailable.")):
            response = self.client.post("/api/chatbot", json=request_body, headers=self.auth())
        self.assertEqual(response.status_code, 503)
        self.assertIn("temporarily unavailable", response.json["error"])

    def test_communication_coach_has_instructions_for_each_mode_and_skips_missing_branch_topics(self):
        response_json = json.dumps({"reply": "What would you like to talk about?", "difficulty": "beginner",
                                    "evaluation": None, "corrections": [], "vocabulary_learned": []})
        expected_prompts = {
            "daily": "natural daily conversation",
            "interview": "branch/course-aware",
            "professional": "realistic workplace roleplays",
            "vocabulary": "Teach one useful word at a time",
            "grammar": "Notice grammar errors",
        }
        for mode, expected in expected_prompts.items():
            body = {"mode": "communication_coach", "practice_mode": mode,
                    "messages": [{"role": "user", "content": "Let's practice."}]}
            with patch("app.request_groq", return_value=(response_json, "groq-test")) as call:
                response = self.client.post("/api/chatbot", json=body, headers=self.auth())
            self.assertEqual(response.status_code, 200, mode)
            self.assertIn(expected, call.call_args.args[1], mode)

        body = {"mode": "communication_coach", "practice_mode": "interview",
                "messages": [{"role": "user", "content": "Ask me an interview question."}]}
        empty_branch_profile = {"branch": "", "year": None, "career_goal": "", "skills": []}
        with patch("app.student_data", return_value=empty_branch_profile):
            with patch("app.request_groq", return_value=(response_json, "groq-test")) as call:
                response = self.client.post("/api/chatbot", json=body, headers=self.auth())
        self.assertEqual(response.status_code, 200)
        self.assertIn('"branch":"not set"', call.call_args.args[1])
        self.assertIn("otherwise stay general", call.call_args.args[1])

    def test_communication_coach_frontend_has_modes_mobile_layout_and_failure_retry(self):
        with open(os.path.join(os.path.dirname(__file__), "../../frontend/assets/js/app.js"), encoding="utf-8") as source_file:
            javascript = source_file.read()
        with open(os.path.join(os.path.dirname(__file__), "../../frontend/assets/css/app.css"), encoding="utf-8") as source_file:
            stylesheet = source_file.read()
        for mode in ("daily", "interview", "professional", "vocabulary", "grammar"):
            self.assertIn(f"{mode}: {{ label:", javascript)
        self.assertIn("data-action=\"end-communication-practice\"", javascript)
        self.assertIn("data-action=\"retry-communication\"", javascript)
        self.assertIn("message.corrected_version", javascript)
        self.assertIn("message.mistake_explanations", javascript)
        self.assertIn("message.vocabulary_suggestions", javascript)
        self.assertIn("@media(max-width:600px)", stylesheet)
        self.assertIn(".coach-report-scores { grid-template-columns:repeat(2", stylesheet)

    def test_missing_groq_key_fails_without_a_local_ai_response(self):
        with patch("ai_service.dotenv_values", return_value={"GROQ_API_KEY": ""}):
            with self.assertRaisesRegex(AIServiceError, "GROQ_API_KEY"):
                request_groq([{"role": "user", "content": "Hello"}], "CAMPUSLINK test context")

    def test_groq_http_statuses_have_distinct_safe_errors(self):
        messages = [{"role": "user", "content": "Hello"}]
        expected = {
            401: "authentication failed",
            403: "access was forbidden",
            404: "model 'groq-test' was not found",
            429: "rate limiting requests",
            500: r"server error \(HTTP 500\)",
            503: "HTTP 503",
        }
        with patch("ai_service.dotenv_values", return_value={"GROQ_API_KEY": "test-key", "GROQ_MODEL": "groq-test"}), \
             patch("ai_service.time.sleep"):
            for status, message in expected.items():
                with self.subTest(status=status):
                    def fail_request(*args, **kwargs):
                        raise HTTPError("https://example.invalid", status, "provider error", None,
                                        BytesIO(b'{"error":{"message":"diagnostic detail"}}'))

                    with patch("ai_service.urlopen", side_effect=fail_request):
                        with self.assertRaisesRegex(AIServiceError, message):
                            request_groq(messages, "CAMPUSLINK test context")

    def test_groq_cloudflare_1010_has_actionable_error(self):
        def blocked_request(*args, **kwargs):
            raise HTTPError("https://example.invalid", 403, "forbidden", None,
                            BytesIO(b"error code: 1010"))

        with patch("ai_service.dotenv_values", return_value={"GROQ_API_KEY": "test-key", "GROQ_MODEL": "groq-test"}), \
             patch("ai_service.urlopen", side_effect=blocked_request):
            with self.assertRaisesRegex(AIServiceError, "edge security blocked.*1010"):
                request_groq([{"role": "user", "content": "Hello"}], "CAMPUSLINK test context")

    def test_groq_retries_temporary_503_before_success(self):
        class SuccessfulResponse:
            status = 200

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return b'{"choices":[{"message":{"content":"Hello from Groq"}}]}'

        responses = [503, 503, SuccessfulResponse()]

        def overloaded_then_success(*args, **kwargs):
            response = responses.pop(0)
            if response == 503:
                raise HTTPError("https://example.invalid", 503, "overloaded", None,
                                BytesIO(b'{"error":{"message":"temporary overload"}}'))
            return response

        with patch("ai_service.dotenv_values", return_value={"GROQ_API_KEY": "test-key", "GROQ_MODEL": "groq-test"}), \
             patch("ai_service.urlopen", side_effect=overloaded_then_success) as send, \
             patch("ai_service.time.sleep") as sleep:
            answer, model = request_groq([{"role": "user", "content": "Hello"}], "CAMPUSLINK test context")

        self.assertEqual((answer, model), ("Hello from Groq", "groq-test"))
        self.assertEqual(send.call_count, 3)
        self.assertEqual(sleep.call_count, 2)
        request = send.call_args.args[0]
        self.assertEqual(request.full_url, GROQ_API_ENDPOINT)
        self.assertEqual(request.get_header("Authorization"), "Bearer test-key")
        self.assertEqual(request.get_header("User-agent"), "CampusLink/1.0 (Groq API client)")
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual(payload["model"], "groq-test")
        self.assertEqual(payload["messages"], [
            {"role": "system", "content": "CAMPUSLINK test context"},
            {"role": "user", "content": "Hello"},
        ])

    def test_schema_upgrade_adds_quiz_metadata_columns_without_losing_questions(self):
        with app.app_context():
            quiz = Quiz(title="Legacy quiz", category="CSE", active=True)
            db.session.add(quiz)
            db.session.flush()
            db.session.add(QuizQuestion(quiz_id=quiz.id, prompt="Keep this question", options='["A","B"]',
                                        answer_index=0, explanation="Existing explanation"))
            db.session.commit()
            for column in ("difficulty", "topic", "question_type", "answer_indices"):
                db.session.execute(text(f"ALTER TABLE quiz_questions DROP COLUMN {column}"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN college_name"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN graduation_year"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN interests"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN gender"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN courses"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN achievements"))
            db.session.execute(text("ALTER TABLE students DROP COLUMN interest_tags"))
            db.session.commit()
            db.session.remove()

            ensure_schema_compatibility()
            saved = QuizQuestion.query.filter_by(prompt="Keep this question").one()
            self.assertEqual(saved.answer_index, 0)
            self.assertEqual(saved.explanation, "Existing explanation")
            self.assertIn(saved.difficulty, {"Easy", "Medium", "Hard"})
            self.assertEqual(saved.topic, "CSE")
            self.assertEqual(saved.question_type, "MCQ")
            self.assertEqual(json.loads(saved.answer_indices or "[]"), [])
            migrated_student = db.session.get(Student, self.student_id)
            self.assertEqual(migrated_student.college_name, "")
            self.assertIsNone(migrated_student.graduation_year)
            self.assertEqual(migrated_student.interests, "")
            self.assertEqual(migrated_student.gender, "")
            self.assertEqual(migrated_student.courses, "[]")
            self.assertEqual(migrated_student.achievements, "[]")
            self.assertEqual(migrated_student.interest_tags, "[]")

    def generated_questions(self, count, category="Python", topic="", difficulty="Medium"):
        items = []
        subject = topic or category
        prompts = [
            "How does {subject} apply its core rule in a normal program?",
            "A student is debugging a {subject} example; which operation is correct?",
            "Choose the accurate principle for this {subject} scenario.",
            "When using {subject}, what result should this method produce?",
            "Which design decision is valid for this {subject} task?",
            "A practice problem asks about {subject}; what is the correct approach?",
            "Identify the behavior expected from this {subject} feature.",
            "For this {subject} case, select the technically sound explanation.",
            "What is the proper way to reason about this {subject} operation?",
            "Which statement best describes this {subject} behavior?",
        ]
        for index in range(count):
            items.append({
                "prompt": prompts[index % len(prompts)].format(subject=subject) + f" Example {index + 1}.",
                "topic": topic or category,
                "difficulty": difficulty,
                "question_type": "MCQ",
                "options": [
                    f"The correct concept {index + 1}",
                    f"An unrelated alternative {index + 1}",
                    f"A misleading rule {index + 1}",
                    f"A reversed idea {index + 1}",
                ],
                "answer_indices": [0],
                "explanation": f"The correct concept {index + 1} describes the rule accurately and distinguishes it from the unrelated alternatives.",
            })
        return items

    def test_quick_quiz_is_ai_generated_once_daily_resumable_and_saved(self):
        generated = self.generated_questions(10, category="Python", topic="decorators")
        with patch("app.generate_ai_quiz_questions", return_value=generated) as generator:
            started = self.client.post("/api/quizzes/quick/start", json={
                "category": "Python", "topic": "decorators", "difficulty": "Medium", "count": 10,
            }, headers=self.auth())
        self.assertEqual(started.status_code, 201, started.json)
        self.assertEqual(generator.call_count, 1)
        self.assertEqual(len(started.json["questions"]), 10)
        self.assertTrue(all(len(question["options"]) == 4 for question in started.json["questions"]))
        self.assertTrue(all("answer_index" not in question and "answer_indices" not in question
                            for question in started.json["questions"]))
        token_payload = quick_quiz_serializer.loads(started.json["quiz"]["attempt_token"])
        self.assertEqual(token_payload["uid"], self.user_id)
        self.assertEqual(token_payload["student_id"], self.student_id)

        options = self.client.get("/api/quizzes/quick/options", headers=self.auth())
        self.assertEqual(options.status_code, 200)
        self.assertEqual(options.json["daily"]["status"], "IN_PROGRESS")
        self.assertIn("Logical Reasoning", options.json["categories"])
        self.assertIn("Custom Topic", options.json["categories"])
        self.assertEqual(len(options.json["active_attempt"]["questions"]), 10)
        self.assertTrue(all("answer_indices" not in question
                            for question in options.json["active_attempt"]["questions"]))

        answers = {str(question["id"]): 0 for question in started.json["questions"]}
        saved = self.client.post("/api/quizzes/quick/progress", json={
            "attempt_token": started.json["quiz"]["attempt_token"], "answers": answers, "revision": 1,
        }, headers=self.auth())
        self.assertEqual(saved.status_code, 200, saved.json)
        resumed = self.client.get("/api/quizzes/quick/options", headers=self.auth())
        self.assertEqual(resumed.json["active_attempt"]["answers"], answers)

        blocked = self.client.post("/api/quizzes/quick/start", json={
            "category": "Python", "difficulty": "Easy", "count": 3,
        }, headers=self.auth())
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(blocked.json["next_eligible_date"],
                         (quick_quiz_local_date() + timedelta(days=1)).isoformat())

        hidden = self.client.post("/api/quizzes/answer-feedback", json={
            "attempt_token": started.json["quiz"]["attempt_token"],
            "question_id": started.json["questions"][0]["id"], "answer": 0,
        }, headers=self.auth())
        self.assertEqual(hidden.status_code, 400)
        self.assertNotIn("correct_indices", hidden.json)

        submitted = self.client.post("/api/quizzes/quick/submit", json={
            "attempt_token": started.json["quiz"]["attempt_token"], "answers": answers,
        }, headers=self.auth())
        self.assertEqual(submitted.status_code, 200, submitted.json)
        self.assertEqual(submitted.json["result"]["score"], 100)
        self.assertEqual(submitted.json["result"]["category"], "Python")
        self.assertTrue(submitted.json["result"]["completed_at"])
        self.assertTrue(all(row["explanation"] for row in submitted.json["result"]["answers"]))
        self.assertEqual(submitted.json["result"]["next_eligible_date"],
                         (quick_quiz_local_date() + timedelta(days=1)).isoformat())
        with app.app_context():
            attempt = QuickQuizAttempt.query.one()
            self.assertEqual(attempt.status, "COMPLETED")
            self.assertIsNotNone(attempt.completed_at)
            self.assertEqual(json.loads(attempt.answers), answers)
            metadata = json.loads(QuizResult.query.one().weak_areas)
            self.assertEqual(len(metadata["answers"]), 10)
            self.assertEqual(metadata["category"], "Python")
            duplicate = QuickQuizAttempt(
                user_id=self.user_id, student_id=self.student_id, attempt_date=attempt.attempt_date,
                attempt_id="another-attempt-id", status="GENERATING", category="Python",
                difficulty="Easy", requested_difficulty="Easy", question_count=3,
            )
            db.session.add(duplicate)
            with self.assertRaises(IntegrityError):
                db.session.commit()
            db.session.rollback()

        replay = self.client.post("/api/quizzes/quick/submit", json={
            "attempt_token": started.json["quiz"]["attempt_token"], "answers": answers,
        }, headers=self.auth())
        self.assertEqual(replay.status_code, 409)
        after_complete = self.client.get("/api/quizzes/quick/options", headers=self.auth())
        self.assertEqual(after_complete.json["daily"]["status"], "COMPLETED")
        self.assertTrue(after_complete.json["daily"]["used"])

    def test_quick_quiz_requires_authentication_and_supported_categories(self):
        unauthenticated = self.client.post("/api/quizzes/quick/start", json={"count": 10})
        self.assertEqual(unauthenticated.status_code, 401)
        invalid = self.client.post("/api/quizzes/quick/start", json={
            "category": "Not A Supported Category", "difficulty": "Medium", "count": 10,
        }, headers=self.auth())
        self.assertEqual(invalid.status_code, 400)
        with app.app_context():
            self.assertEqual(QuickQuizAttempt.query.count(), 0)

    def test_custom_topic_requires_a_topic_and_generates_that_subject(self):
        missing_topic = self.client.post("/api/quizzes/quick/start", json={
            "category": "Custom Topic", "difficulty": "Easy", "count": 3,
        }, headers=self.auth())
        self.assertEqual(missing_topic.status_code, 400)
        generated = self.generated_questions(3, category="Custom Topic", topic="REST API design", difficulty="Easy")
        with patch("app.generate_ai_quiz_questions", return_value=generated):
            response = self.client.post("/api/quizzes/quick/start", json={
                "category": "Custom Topic", "topic": "REST API design", "difficulty": "Easy", "count": 3,
            }, headers=self.auth())
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json["quiz"]["category"], "Custom Topic")
        self.assertEqual(response.json["quiz"]["topic"], "REST API design")
        self.assertTrue(all(question["topic"] == "REST API design" for question in response.json["questions"]))

    def test_ai_topic_quiz_is_validated_and_multiple_answers_are_graded_after_submit(self):
        generated = self.generated_questions(3, category="Data Structures", topic="Linked Lists")
        generated[1]["question_type"] = "MULTIPLE_ANSWER"
        generated[1]["answer_indices"] = [0, 2]
        generated[1]["explanation"] = "The correct concept and misleading rule are both required for this answer."
        generated[2]["question_type"] = "TRUE_FALSE"
        generated[2]["options"] = ["True", "False"]
        generated[2]["answer_indices"] = [0]
        generated[2]["explanation"] = "True is correct because the statement accurately reflects the linked list behavior."
        duplicate = {
            "prompt": "Data Structures test question 1?", "topic": "Linked Lists",
            "difficulty": "Medium", "question_type": "MCQ", "options": ["One", "Two"],
            "answer_indices": [0], "explanation": "The first response repeats an existing question.",
        }
        provider = patch("app.request_quiz_generation", side_effect=[
            (json.dumps({"questions": [duplicate, duplicate, duplicate]}), "test-model", "OpenRouter"),
            (json.dumps({"questions": generated}), "test-model", "OpenRouter"),
        ])
        with provider as mock_provider:
            response = self.client.post("/api/quizzes/quick/start", json={
                "category": "Data Structures", "topic": "Linked Lists", "difficulty": "Medium", "count": 3,
            }, headers=self.auth())
        self.assertEqual(mock_provider.call_count, 2)
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json["quiz"]["ai_generated_count"], 3)
        self.assertTrue(all(question["topic"] == "Linked Lists" for question in response.json["questions"]))
        multiple = next(question for question in response.json["questions"]
                        if question["question_type"] == "MULTIPLE_ANSWER")
        feedback = self.client.post("/api/quizzes/answer-feedback", json={
            "attempt_token": response.json["quiz"]["attempt_token"],
            "question_id": multiple["id"], "answer": [0],
        }, headers=self.auth())
        self.assertEqual(feedback.status_code, 400)
        self.assertNotIn("correct_options", feedback.json)

        correct = {}
        with app.app_context():
            for question in response.json["questions"]:
                row = db.session.get(QuizQuestion, question["id"])
                correct[str(row.id)] = json.loads(row.answer_indices or "[]") or [row.answer_index]
        correct[str(multiple["id"])] = correct[str(multiple["id"])][:1]
        submitted = self.client.post("/api/quizzes/quick/submit", json={
            "attempt_token": response.json["quiz"]["attempt_token"], "answers": correct,
        }, headers=self.auth())
        self.assertEqual(submitted.status_code, 200, submitted.json)
        self.assertEqual(submitted.json["result"]["correct"], 2)
        self.assertEqual(submitted.json["result"]["total"], 3)
        self.assertTrue(submitted.json["result"]["topic_scores"])

    def test_ai_generation_failure_releases_daily_slot_and_does_not_fall_back_to_local_questions(self):
        self.add_question_pool(["Python"], per_category=5)
        generated = self.generated_questions(3, category="Python", topic="classes")
        with patch("app.generate_ai_quiz_questions", side_effect=[
            AIServiceError("service unavailable"), generated,
        ]) as generator:
            failed = self.client.post("/api/quizzes/quick/start", json={
                "category": "Python", "topic": "classes", "difficulty": "Medium", "count": 3,
            }, headers=self.auth())
            self.assertEqual(failed.status_code, 503, failed.json)
            self.assertFalse(failed.json["daily_allowance_used"])
            with app.app_context():
                self.assertEqual(QuickQuizAttempt.query.count(), 0)
                self.assertEqual(QuizResult.query.count(), 0)
            retried = self.client.post("/api/quizzes/quick/start", json={
                "category": "Python", "topic": "classes", "difficulty": "Medium", "count": 3,
            }, headers=self.auth())
        self.assertEqual(retried.status_code, 201, retried.json)
        self.assertEqual(generator.call_count, 2)
        self.assertEqual(len(retried.json["questions"]), 3)

    def test_quick_quiz_daily_date_uses_configured_timezone(self):
        class FixedDateTime(datetime):
            @classmethod
            def now(cls, tz=None):
                fixed = cls(2026, 1, 1, 2, 0, 0)
                return fixed.replace(tzinfo=timezone.utc) if tz else fixed

        previous = app.config.get("APP_TIMEZONE")
        with app.app_context(), patch("app.datetime", FixedDateTime):
            app.config["APP_TIMEZONE"] = "Asia/Kolkata"
            self.assertEqual(quick_quiz_local_date(), date(2026, 1, 1))
            app.config["APP_TIMEZONE"] = "America/Los_Angeles"
            self.assertEqual(quick_quiz_local_date(), date(2025, 12, 31))
        app.config["APP_TIMEZONE"] = previous

    def test_openrouter_quiz_request_uses_configured_model_and_retries_temporary_errors(self):
        class SuccessfulResponse:
            status = 200

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self):
                return b'{"choices":[{"message":{"content":"{\\"questions\\":[]}"}}]}'

        responses = [503, 503, SuccessfulResponse()]

        def overloaded_then_success(*args, **kwargs):
            response = responses.pop(0)
            if response == 503:
                raise HTTPError("https://example.invalid", 503, "overloaded", None,
                                BytesIO(b'{"error":{"message":"temporary overload"}}'))
            return response

        messages = [{"role": "user", "content": "Generate a Python quiz"}]
        with patch("ai_service.dotenv_values", return_value={
            "OPENROUTER_API_KEY": "test-openrouter-key",
            "QUIZ_MODEL": "deepseek/deepseek-v4.1-flash",
        }), patch("ai_service.urlopen", side_effect=overloaded_then_success) as send, \
                patch("ai_service.time.sleep") as sleep:
            answer, model, provider = request_quiz_generation(messages, "Return JSON")

        self.assertEqual((answer, model, provider),
                         ('{"questions":[]}', "deepseek/deepseek-v4.1-flash", "OpenRouter"))
        self.assertEqual(send.call_count, 3)
        self.assertEqual(sleep.call_count, 2)
        request = send.call_args.args[0]
        self.assertEqual(request.full_url, OPENROUTER_CHAT_ENDPOINT)
        self.assertEqual(request.get_header("Authorization"), "Bearer test-openrouter-key")
        self.assertEqual(json.loads(request.data.decode("utf-8"))["model"], "deepseek/deepseek-v4.1-flash")

    def test_openrouter_configuration_and_http_errors_are_safe(self):
        with patch("ai_service.dotenv_values", return_value={"OPENROUTER_API_KEY": ""}):
            with self.assertRaisesRegex(AIServiceError, "OPENROUTER_API_KEY"):
                request_quiz_generation([{"role": "user", "content": "Hello"}], "Quiz")

        messages = [{"role": "user", "content": "Generate a quiz"}]
        expected = {
            401: "authentication failed",
            403: "access was forbidden",
            404: "could not find quiz model",
            402: "requires available credits",
            429: "rate limiting",
            500: "temporarily unavailable",
            503: "temporarily unavailable",
        }
        with patch("ai_service.dotenv_values", return_value={
            "OPENROUTER_API_KEY": "private-test-key", "QUIZ_MODEL": "deepseek/deepseek-v4.1-flash",
        }), patch("ai_service.time.sleep"):
            for status, message in expected.items():
                with self.subTest(status=status):
                    def fail_request(*args, **kwargs):
                        raise HTTPError("https://example.invalid", status, "provider error", None,
                                        BytesIO(b'{"error":{"message":"diagnostic detail"}}'))

                    with patch("ai_service.urlopen", side_effect=fail_request):
                        with self.assertRaisesRegex(AIServiceError, message) as raised:
                            request_quiz_generation(messages, "Quiz")
                    self.assertNotIn("private-test-key", str(raised.exception))
    def test_adaptive_quiz_difficulty_moves_from_recent_performance(self):
        with app.app_context():
            quiz = Quiz(title="Adaptive history", category="CSE", active=True)
            db.session.add(quiz)
            db.session.flush()
            for index, (score, level) in enumerate(((90, "Medium"), (88, "Medium"), (85, "Medium"))):
                metadata = {"type": "full_quiz", "category": "Mixed", "difficulty": level}
                db.session.add(QuizResult(student_id=self.student_id, quiz_id=quiz.id, score=score,
                                          correct_count=9, weak_areas=json.dumps(metadata),
                                          created_at=datetime.utcnow() - timedelta(days=2-index)))
            db.session.commit()
            self.assertEqual(adaptive_quiz_difficulty(self.student_id), "Hard")

    def test_full_quiz_attempt_is_branch_scoped_balanced_and_single_use(self):
        self.add_question_pool(["Python"], per_category=6)
        with app.app_context():
            quiz = Quiz.query.filter_by(category="Python").first()
            questions = QuizQuestion.query.filter_by(quiz_id=quiz.id).order_by(QuizQuestion.id).all()
            questions[0].difficulty = "Easy"
            questions[1].difficulty = "Hard"
            questions[0].question_type = "MULTIPLE_ANSWER"
            questions[0].answer_indices = "[0,2]"
            db.session.commit()
            quiz_id = quiz.id
        start = self.client.post(f"/api/quizzes/{quiz_id}/start", json={}, headers=self.auth())
        self.assertEqual(start.status_code, 201, start.json)
        payload = full_quiz_serializer.loads(start.json["quiz"]["attempt_token"])
        self.assertEqual(payload["uid"], self.user_id)
        self.assertEqual(payload["student_id"], self.student_id)
        self.assertTrue(all("answer_index" not in row and "answer_indices" not in row for row in start.json["questions"]))
        feedback = self.client.post("/api/quizzes/answer-feedback", json={
            "attempt_token": start.json["quiz"]["attempt_token"],
            "question_id": start.json["questions"][0]["id"], "answer": [0, 2],
        }, headers=self.auth())
        self.assertEqual(feedback.status_code, 200, feedback.json)
        self.assertIn("explanation", feedback.json["feedback"])
        levels = [row["difficulty"] for row in start.json["questions"]]
        self.assertEqual(levels, sorted(levels, key={"Easy": 0, "Medium": 1, "Hard": 2}.get))
        answers = {}
        with app.app_context():
            for question in start.json["questions"]:
                row = db.session.get(QuizQuestion, question["id"])
                answers[str(row.id)] = json.loads(row.answer_indices or "[]") or [row.answer_index]
        submit = self.client.post(f"/api/quizzes/{quiz_id}/submit", json={
            "attempt_token": start.json["quiz"]["attempt_token"], "answers": answers,
        }, headers=self.auth())
        self.assertEqual(submit.status_code, 200, submit.json)
        self.assertEqual(submit.json["result"]["score"], 100)
        replay = self.client.post(f"/api/quizzes/{quiz_id}/submit", json={
            "attempt_token": start.json["quiz"]["attempt_token"], "answers": answers,
        }, headers=self.auth())
        self.assertEqual(replay.status_code, 400)
        repeated = self.client.post(f"/api/quizzes/{quiz_id}/start", json={}, headers=self.auth())
        self.assertEqual(repeated.status_code, 201)
        self.assertTrue(repeated.json["quiz"]["reused_questions"])

    def test_quiz_frontend_supports_topic_adaptive_types_refresh_and_mobile(self):
        root = Path(__file__).resolve().parents[2]
        script = (root / "frontend" / "assets" / "js" / "app.js").read_text(encoding="utf-8")
        styles = (root / "frontend" / "assets" / "css" / "app.css").read_text(encoding="utf-8")
        self.assertIn("name=\"topic\"", script)
        self.assertIn("value=\"Adaptive\"", script)
        self.assertIn("question.question_type==='MULTIPLE_ANSWER'", script)
        self.assertIn("/quizzes/answer-feedback", script)
        self.assertIn("role=\"progressbar\"", script)
        self.assertIn("sessionStorage.setItem('campuslink_quiz_attempt'", script)
        self.assertIn("sessionStorage.getItem('campuslink_quiz_attempt'", script)
        self.assertIn("/quizzes/${id}/start", script)
        self.assertIn(".quick-quiz-controls { grid-template-columns:1fr", styles)


if __name__ == "__main__":
    unittest.main()
